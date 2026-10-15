#!/usr/bin/env python3
"""doc_read.py — read pdf/docx/xlsx/csv/md/txt/json into text or JSON.

  doc_read.py FILE [--max-chars N] [--pages A-B] [--max-rows N] [--json]
  doc_read.py --selftest
"""
from __future__ import annotations

import csv
import json
import os
import sys
import tempfile
from pathlib import Path


def read_pdf(p: Path, pages: str | None) -> dict:
    from pypdf import PdfReader  # pip install pypdf
    r = PdfReader(str(p))
    total = len(r.pages)
    lo, hi = 1, total
    if pages:
        a, _, b = pages.partition("-")
        lo, hi = int(a), int(b or a)
    out, empty = [], 0
    for i in range(lo - 1, min(hi, total)):
        t = (r.pages[i].extract_text() or "").strip()
        if not t:
            empty += 1
        out.append(f"--- page {i + 1} ---\n{t}")
    note = "no text layer on some pages (scanned?) — OCR needed: tesseract + pytesseract" if empty else ""
    return {"type": "pdf", "pages": total, "text": "\n".join(out), "note": note}


def read_docx(p: Path) -> dict:
    import docx  # pip install python-docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    tables = []
    for t in d.tables:
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        tables.append(rows)
        parts.append("\n".join(" | ".join(r) for r in rows))
    return {"type": "docx", "paragraphs": len(d.paragraphs), "tables": len(tables), "text": "\n".join(parts)}


def read_xlsx(p: Path, max_rows: int) -> dict:
    import openpyxl  # pip install openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    sheets, parts = {}, []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(["" if v is None else v for v in row])
        sheets[ws.title] = rows
        parts.append(f"--- sheet {ws.title} ({ws.max_row} rows x {ws.max_column} cols) ---")
        parts += [" | ".join(str(c) for c in r) for r in rows]
    return {"type": "xlsx", "sheets": list(sheets), "rows": sheets, "text": "\n".join(parts)}


def read_csv(p: Path, max_rows: int) -> dict:
    delim = "\t" if p.suffix.lower() == ".tsv" else ","
    with p.open(encoding="utf-8", errors="replace", newline="") as f:
        rows = []
        for i, r in enumerate(csv.reader(f, delimiter=delim)):
            if i >= max_rows:
                break
            rows.append(r)
    return {"type": "csv", "rows": rows, "count": len(rows), "text": "\n".join(" | ".join(r) for r in rows)}


def read_any(path: str, pages: str | None = None, max_rows: int = 200) -> dict:
    p = Path(path)
    if not p.is_file():
        return {"error": f"not a file: {path}"}
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            res = read_pdf(p, pages)
        elif ext == ".docx":
            res = read_docx(p)
        elif ext in (".xlsx", ".xlsm"):
            res = read_xlsx(p, max_rows)
        elif ext in (".csv", ".tsv"):
            res = read_csv(p, max_rows)
        elif ext == ".json":
            data = json.loads(p.read_text(encoding="utf-8"))
            res = {"type": "json", "text": json.dumps(data, ensure_ascii=False, indent=1)[:200000]}
        else:
            res = {"type": "text", "text": p.read_text(encoding="utf-8", errors="replace")}
    except ImportError as e:
        return {"error": f"missing library: {e.name} (pip install pypdf python-docx openpyxl)"}
    except Exception as e:
        return {"error": f"{type(e).__name__}: {str(e)[:200]}"}
    res["file"] = str(p)
    return res


def selftest() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="dd-doc-"))
    c = tmp / "a.csv"
    c.write_text("name,amount\nx,10\ny,20\n", encoding="utf-8")
    r = read_any(str(c))
    assert r["count"] == 3 and r["rows"][1] == ["x", "10"], r
    try:
        import docx
        d = docx.Document()
        d.add_paragraph("Hello dd")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "k", "v"
        d.save(str(tmp / "a.docx"))
        rd = read_any(str(tmp / "a.docx"))
        assert "Hello dd" in rd["text"] and "k | v" in rd["text"], rd
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["h1", "h2"])
        wb.active.append([1, 2])
        wb.save(str(tmp / "a.xlsx"))
        rx = read_any(str(tmp / "a.xlsx"))
        assert rx["rows"]["Sheet"][1] == [1, 2], rx
        from reportlab.pdfgen import canvas
        cv = canvas.Canvas(str(tmp / "a.pdf"))
        cv.drawString(72, 720, "pdf text dd")
        cv.save()
        rp = read_any(str(tmp / "a.pdf"))
        assert "pdf text dd" in rp["text"], rp
        print("doc_read selftest ok (csv, docx, xlsx, pdf)")
    except ImportError as e:
        print(f"doc_read selftest ok (csv only; install {e.name} for more)")
    return 0


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    if argv[0] == "--selftest":
        return selftest()
    mc = int(argv[argv.index("--max-chars") + 1]) if "--max-chars" in argv else 20000
    pages = argv[argv.index("--pages") + 1] if "--pages" in argv else None
    mr = int(argv[argv.index("--max-rows") + 1]) if "--max-rows" in argv else 200
    res = read_any(argv[0], pages, mr)
    if "error" in res:
        print(json.dumps(res))
        return 1
    if "--json" in argv:
        res["text"] = res["text"][:mc]
        print(json.dumps(res, ensure_ascii=False, indent=1, default=str))
    else:
        head = f"[{res['type']}] {res['file']}" + (f" — {res['note']}" if res.get("note") else "")
        print(head + "\n" + res["text"][:mc])
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
