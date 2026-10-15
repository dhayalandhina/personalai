#!/usr/bin/env python3
"""doc_write.py — build docx / pdf / xlsx / csv from Markdown, CSV or JSON rows.

  doc_write.py docx OUT --md FILE.md [--title T]
  doc_write.py pdf  OUT --md FILE.md [--title T]
  doc_write.py xlsx OUT (--csv FILE | --json rows.json) [--sheet NAME]
  doc_write.py csv  OUT --json rows.json
  doc_write.py --selftest
Markdown subset: # ## ### headings, paragraphs, "- " bullets, "1. " numbered, **bold**, `code`, | tables |, --- rule.
"""
from __future__ import annotations

import csv
import json
import re
import sys
import tempfile
from pathlib import Path


# ---------- markdown mini-parser ----------
def parse_md(text: str) -> list[dict]:
    blocks: list[dict] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            blocks.append({"t": "h", "level": len(m.group(1)), "text": m.group(2)})
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append({"t": "table", "rows": rows})
            continue
        elif re.match(r"^\s*[-*]\s+", ln):
            blocks.append({"t": "li", "text": re.sub(r"^\s*[-*]\s+", "", ln)})
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            blocks.append({"t": "num", "text": re.sub(r"^\s*\d+[.)]\s+", "", ln)})
        elif ln.strip() == "---":
            blocks.append({"t": "hr"})
        else:
            para = [ln.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#{1,3}\s|\||\s*[-*]\s|\s*\d+[.)]\s|---$)", lines[i + 1]):
                i += 1
                para.append(lines[i].strip())
            blocks.append({"t": "p", "text": " ".join(para)})
        i += 1
    return blocks


def runs(text: str) -> list[tuple[str, bool, bool]]:
    """split **bold** and `code` into (text, bold, code) runs"""
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("`"):
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


# ---------- docx ----------
def write_docx(out: str, md: str, title: str | None) -> dict:
    import docx  # pip install python-docx
    from docx.shared import Pt
    d = docx.Document()
    if title:
        d.add_heading(title, 0)
    for b in parse_md(md):
        if b["t"] == "h":
            d.add_heading(b["text"], b["level"])
        elif b["t"] in ("p", "li", "num"):
            style = {"p": None, "li": "List Bullet", "num": "List Number"}[b["t"]]
            para = d.add_paragraph(style=style) if style else d.add_paragraph()
            for txt, bold, code in runs(b["text"]):
                r = para.add_run(txt)
                r.bold = bold
                if code:
                    r.font.name = "Courier New"
                    r.font.size = Pt(10)
        elif b["t"] == "table" and b["rows"]:
            t = d.add_table(rows=len(b["rows"]), cols=max(len(r) for r in b["rows"]))
            t.style = "Table Grid"
            for ri, row in enumerate(b["rows"]):
                for ci, cell in enumerate(row):
                    t.cell(ri, ci).text = cell
                    if ri == 0:
                        for r in t.cell(ri, ci).paragraphs[0].runs:
                            r.bold = True
        elif b["t"] == "hr":
            d.add_paragraph("_" * 40)
    d.save(out)
    return {"ok": True, "path": out, "blocks": len(parse_md(md))}


# ---------- pdf ----------
def write_pdf(out: str, md: str, title: str | None) -> dict:
    from reportlab.lib import colors  # pip install reportlab
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, ListFlowable, ListItem
    ss = getSampleStyleSheet()
    story = []
    if title:
        story += [Paragraph(title, ss["Title"]), Spacer(1, 0.3 * cm)]

    def rl(text: str) -> str:
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
        return re.sub(r"`([^`]+)`", r"<font face='Courier'>\1</font>", text)

    pending: list = []
    kind = None

    def flush():
        nonlocal pending, kind
        if pending:
            story.append(ListFlowable([ListItem(Paragraph(rl(t), ss["BodyText"])) for t in pending],
                                      bulletType="1" if kind == "num" else "bullet"))
            story.append(Spacer(1, 0.15 * cm))
        pending, kind = [], None

    for b in parse_md(md):
        if b["t"] in ("li", "num"):
            if kind and kind != b["t"]:
                flush()
            kind = b["t"]
            pending.append(b["text"])
            continue
        flush()
        if b["t"] == "h":
            story.append(Paragraph(rl(b["text"]), ss[f"Heading{b['level']}"]))
        elif b["t"] == "p":
            story += [Paragraph(rl(b["text"]), ss["BodyText"]), Spacer(1, 0.15 * cm)]
        elif b["t"] == "table" and b["rows"]:
            t = Table(b["rows"], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
            story += [t, Spacer(1, 0.3 * cm)]
        elif b["t"] == "hr":
            story.append(Spacer(1, 0.5 * cm))
    flush()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm).build(story)
    return {"ok": True, "path": out}


# ---------- xlsx / csv ----------
def rows_from(csv_path: str | None, json_path: str | None) -> list[list]:
    if csv_path:
        with open(csv_path, encoding="utf-8", newline="") as f:
            return [r for r in csv.reader(f)]
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    if data and isinstance(data[0], dict):
        keys = list(data[0].keys())
        return [keys] + [[row.get(k, "") for k in keys] for row in data]
    return data


def write_xlsx(out: str, rows: list[list], sheet: str = "Sheet1") -> dict:
    import openpyxl  # pip install openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for r in rows:
        ws.append([_num(c) for c in r])
    if rows:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EEF5")
        ws.freeze_panes = "A2"
        for i in range(1, len(rows[0]) + 1):
            width = max(len(str(r[i - 1])) if i - 1 < len(r) else 0 for r in rows)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    wb.save(out)
    return {"ok": True, "path": out, "rows": len(rows), "sheet": ws.title}


def _num(v):
    if isinstance(v, str):
        s = v.replace(",", "")
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return v


def write_csv(out: str, rows: list[list]) -> dict:
    with open(out, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)
    return {"ok": True, "path": out, "rows": len(rows)}


MD_FIXTURE = """# Quotation
Client: **Aqua Ltd** — ref `Q-101`
## Items
| Item | Qty | Price |
|---|---|---|
| Purifier | 2 | 12,500 |
- delivery in 7 days
- GST 18% extra
1. Advance 50%
2. Balance on delivery
---
Thanks."""


def selftest() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="dd-docw-"))
    b = parse_md(MD_FIXTURE)
    kinds = [x["t"] for x in b]
    assert kinds == ["h", "p", "h", "table", "li", "li", "num", "num", "hr", "p"], kinds
    assert b[3]["rows"][1] == ["Purifier", "2", "12,500"]
    (tmp / "r.json").write_text(json.dumps([{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]))
    rows = rows_from(None, str(tmp / "r.json"))
    assert rows == [["a", "b"], [1, "x"], [2, "y"]]
    assert write_csv(str(tmp / "o.csv"), rows)["rows"] == 3
    done = ["csv"]
    try:
        write_docx(str(tmp / "o.docx"), MD_FIXTURE, "Test")
        done.append("docx")
        write_xlsx(str(tmp / "o.xlsx"), rows)
        done.append("xlsx")
        write_pdf(str(tmp / "o.pdf"), MD_FIXTURE, "Test")
        done.append("pdf")
        assert (tmp / "o.pdf").stat().st_size > 1000
    except ImportError as e:
        print(f"(install {e.name} for more formats)")
    print(f"doc_write selftest ok ({', '.join(done)})")
    return 0


def arg(argv: list[str], flag: str, default=None):
    return argv[argv.index(flag) + 1] if flag in argv else default


def main(argv: list[str]) -> int:
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    if argv[0] == "--selftest":
        return selftest()
    kind, out = argv[0], argv[1]
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    try:
        if kind in ("docx", "pdf"):
            md = Path(arg(argv, "--md")).read_text(encoding="utf-8")
            res = (write_docx if kind == "docx" else write_pdf)(out, md, arg(argv, "--title"))
        elif kind == "xlsx":
            res = write_xlsx(out, rows_from(arg(argv, "--csv"), arg(argv, "--json")), arg(argv, "--sheet", "Sheet1"))
        elif kind == "csv":
            res = write_csv(out, rows_from(None, arg(argv, "--json")))
        else:
            print(__doc__)
            return 2
    except ImportError as e:
        res = {"error": f"missing library {e.name}: pip install python-docx openpyxl reportlab"}
    except Exception as e:
        res = {"error": f"{type(e).__name__}: {str(e)[:200]}"}
    print(json.dumps(res, ensure_ascii=False))
    return 1 if "error" in res else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
